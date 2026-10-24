"""Build docs/deck/thermal-trace-pitch.pptx from docs/deck/index.html.

Every slide is a full-bleed background picture plus one transparent picture per
animated element of the HTML deck. Each layer gets a native PowerPoint entrance
animation (fade, fly, wipe) with the same start delay, and each slide a fade
transition and its speaker notes. Text inside the pictures is not editable:
edit the HTML (the source of truth) and rebuild.

    python3 scripts/deck_pptx/build.py          # needs: python-pptx, pillow,
                                                #  node + playwright (global)
"""
import html
import json
import os
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Emu

HERE = Path(__file__).resolve().parent
OUT = HERE.parents[1] / "docs" / "deck" / "thermal-trace-pitch.pptx"
SCALE = 1.5  # capture device scale factor (capture.js)
W_EMU, H_EMU = 12192000, 6858000
NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
EMU_PX = W_EMU / 1920


def capture(layers: Path) -> None:
    env = dict(os.environ)
    root = subprocess.run(["npm", "root", "-g"], capture_output=True, text=True).stdout.strip()
    env["NODE_PATH"] = root
    subprocess.run(["node", str(HERE / "capture.js"), str(layers)], check=True, env=env)


class Ids:
    def __init__(self):
        self.n = 2

    def __call__(self):
        self.n += 1
        return self.n


def set_visible(ids, spid):
    return (
        f'<p:set><p:cBhvr><p:cTn id="{ids()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/>'
        f'</p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst>'
        f"<p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>"
        f'<p:to><p:strVal val="visible"/></p:to></p:set>'
    )


def move(ids, spid, attr, frm, dur):
    return (
        f'<p:anim calcmode="lin" valueType="num"><p:cBhvr additive="base"><p:cTn id="{ids()}" dur="{dur}" fill="hold"/>'
        f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>{attr}</p:attrName></p:attrNameLst></p:cBhvr>'
        f'<p:tavLst><p:tav tm="0"><p:val><p:strVal val="{frm}"/></p:val></p:tav>'
        f'<p:tav tm="100000"><p:val><p:strVal val="#{attr}"/></p:val></p:tav></p:tavLst></p:anim>'
    )


def effect_filter(ids, spid, flt, dur):
    return (
        f'<p:animEffect transition="in" filter="{flt}"><p:cBhvr><p:cTn id="{ids()}" dur="{dur}"/>'
        f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
    )


def effect(ids, spid, kind, delay, dur):
    """One entrance effect, started `delay` ms after the slide appears."""
    # (presetID, presetSubtype, default duration ms, body builder)
    if kind in ("draw", "wipe", "grow"):
        pid, sub = 22, 8
        d = dur or {"draw": 1600, "wipe": 1800, "grow": 1400}[kind]
        body = effect_filter(ids, spid, "wipe(left)", d)
    elif kind == "rise":
        pid, sub, d = 22, 4, 2000
        body = effect_filter(ids, spid, "wipe(down)", d)
    elif kind == "left":
        pid, sub, d = 2, 8, 900
        body = move(ids, spid, "ppt_x", "#ppt_x-0.025", d) + effect_filter(ids, spid, "fade", d)
    elif kind == "right":
        pid, sub, d = 2, 2, 900
        body = move(ids, spid, "ppt_x", "#ppt_x+0.025", d) + effect_filter(ids, spid, "fade", d)
    elif kind == "zoom":
        pid, sub, d = 53, 16, 900
        body = (
            move(ids, spid, "ppt_w", "#ppt_w*0.94", d)
            + move(ids, spid, "ppt_h", "#ppt_h*0.94", d)
            + effect_filter(ids, spid, "fade", d)
        )
    elif kind == "fade":
        pid, sub, d = 10, 0, 900
        body = effect_filter(ids, spid, "fade", d)
    else:  # "up": the HTML default, a short rise while fading in
        pid, sub, d = 2, 4, 900
        body = move(ids, spid, "ppt_y", "#ppt_y+0.03", d) + effect_filter(ids, spid, "fade", d)
    return (
        f'<p:par><p:cTn id="{ids()}" presetID="{pid}" presetClass="entr" presetSubtype="{sub}" fill="hold" '
        f'decel="100000" nodeType="withEffect"><p:stCondLst><p:cond delay="{int(delay)}"/></p:stCondLst><p:childTnLst>'
        f"{set_visible(ids, spid)}{body}</p:childTnLst></p:cTn></p:par>"
    )


def timing(effects_xml):
    return etree.fromstring(
        f'<p:timing xmlns:p="{NS}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">'
        f'<p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
        f'<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/>'
        f'<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
        f'<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
        f"{effects_xml}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>"
        f'</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>'
        f'</p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>'
        f"</p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>"
    )


def notes_text(raw: str) -> str:
    t = re.sub(r"</p>\s*", "\n\n", raw)
    t = re.sub(r"<br\s*/?>", "\n", t)
    t = re.sub(r"<[^>]+>", "", t)
    return html.unescape(t).strip()


def main():
    layers = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(tempfile.gettempdir()) / "deck-layers"
    if not (layers / "manifest.json").exists():
        capture(layers)
    manifest = json.loads((layers / "manifest.json").read_text())

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W_EMU), Emu(H_EMU)
    prs.core_properties.title = "Thermal Trace — Pitch"
    prs.core_properties.author = "Generative Gains"
    blank = prs.slide_layouts[6]

    for i, info in enumerate(manifest, 1):
        d = layers / f"{i:02d}"
        slide = prs.slides.add_slide(blank)
        bg = slide.shapes.add_picture(str(d / "bg.png"), 0, 0, Emu(W_EMU), Emu(H_EMU))
        bg.name = "Background"
        bg._element.nvPicPr.cNvPr.set("descr", f"{info['title']} slide background")
        ids, fx = Ids(), []
        for L in info["layers"]:
            pic = slide.shapes.add_picture(
                str(d / L["file"]),
                Emu(round(L["x"] * EMU_PX)), Emu(round(L["y"] * EMU_PX)),
                Emu(round(L["w"] * EMU_PX)), Emu(round(L["h"] * EMU_PX)),
            )
            pic.name = f"{L['id']:02d} {L['kind']} · {L['alt'][:30] or 'graphic'}"
            pic._element.nvPicPr.cNvPr.set("descr", L["alt"])
            fx.append(effect(ids, pic.shape_id, L["kind"], L["delay"], L["dur"]))
        sld = slide._element
        trans = etree.fromstring(f'<p:transition xmlns:p="{NS}" spd="med"><p:fade/></p:transition>')
        sld.append(trans)
        if fx:
            sld.append(timing("".join(fx)))
        notes = notes_text(info["notes"])
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        print(f"slide {i:02d} {info['title']}: {len(fx)} animated layers")

    prs.save(OUT)
    print(f"wrote {OUT} ({OUT.stat().st_size / 1e6:.1f} MB)")


if __name__ == "__main__":
    main()
